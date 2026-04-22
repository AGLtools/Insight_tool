import sys
sys.path.insert(0, r'c:\AGL Analytics\venv_agl\Lib\site-packages')
from pptx import Presentation
try:
    prs = Presentation(r'c:\AGL Analytics\DATA\template.pptx')
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        print(f"\nSlide {i}:")
        for shape in slide.shapes:
            try:
                if hasattr(shape, 'text') and shape.text.strip():
                    print(f"  Shape '{shape.name}' (type={shape.shape_type}): '{shape.text[:100]}'")
            except:
                pass
except Exception as e:
    print(f"Error: {e}")
