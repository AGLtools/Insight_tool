import sys
sys.path.insert(0, r'c:\AGL Analytics\venv_agl\Lib\site-packages')
from pptx import Presentation
try:
    path = r"c:\AGL Analytics\DATA\REPORT\BASE MARS YTD 2026\AGL_Rapport_YTD_Mars_2026.pptx"
    prs3 = Presentation(path)
    print(f"Total slides: {len(prs3.slides)}")
    for i, slide in enumerate(prs3.slides):
        print(f"\nSlide {i}:")
        for shape in slide.shapes:
            try:
                if hasattr(shape, 'text') and shape.text.strip():
                    print(f"  Shape '{shape.name}' (type={shape.shape_type}): '{shape.text[:80]}'")
            except:
                pass
except Exception as e:
    print(f"Error: {e}")
